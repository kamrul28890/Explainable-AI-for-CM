"""Build the formatted Word proposal from the Markdown source.

The script implements the document's page geometry, typography, table XML,
headers/footers, and a small Markdown parser. It intentionally targets the
known proposal structure rather than attempting to be a general Markdown to
DOCX converter.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "Research-prposal.md"
OUTPUT = BASE / "Research-prposal.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 20, 20)
MUTED = RGBColor(85, 85, 85)
TABLE_HEADER = "F4F6F9"
BORDER = "D0D7DE"


def set_cell_shading(cell, fill):
    """Apply a hexadecimal background fill to a python-docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    """Set cell padding in Word twips through the underlying OOXML."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    # python-docx does not expose table-cell margins directly, so create or
    # update each w:tcMar child in the cell properties.
    for margin_name, margin_value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="4"):
    """Create consistent outer and inner borders for a Word table."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    """Set total, grid-column, and individual cell widths in Word twips.

    Updating all three representations prevents Word from discarding the
    intended layout when it recalculates an autofit table.
    """
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    # Repeat widths at cell level because Word readers may prioritize tcW over
    # the table grid depending on compatibility settings.
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths[idx] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))


def add_page_number(paragraph):
    """Insert a live PAGE field into a footer paragraph."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def configure_styles(doc):
    """Configure the document's global body, heading, and list typography."""
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name in ("Title", "Subtitle"):
        style = styles[style_name]
        style.font.name = "Calibri"

    title = styles["Title"]
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = styles["Subtitle"]
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BLUE
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = BLUE
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = DARK_BLUE
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def configure_page(doc):
    """Set margins and install the recurring header and page-number footer."""
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.text = "Research Proposal | Explainable VLMs for Construction Safety"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if p.runs:
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = MUTED

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)
    if p.runs:
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = MUTED


def add_title_block(doc, title):
    """Add the centered title, subtitle, and shaded strategic-focus callout."""
    p = doc.add_paragraph(style="Title")
    p.add_run(title)

    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("Prepared for research discussion with Professor Mustafa Abdallah")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F9")
    p_pr.append(shd)
    run = p.add_run("Strategic focus: ")
    run.bold = True
    run.font.color.rgb = DARK_BLUE
    p.add_run(
        "Dynamic construction safety monitoring using ConstructionSite 10k, SODA, and CMA, "
        "with Florence-2 as the first pilot model."
    )


def add_markdown_table(doc, rows):
    """Render parsed Markdown table rows as a styled Word table."""
    headers = rows[0]
    body = rows[2:] if len(rows) > 1 and set(rows[1]) == {"---"} else rows[1:]
    col_count = len(headers)
    table = doc.add_table(rows=1, cols=col_count)
    table.style = "Table Grid"

    if col_count == 3:
        widths = [2200, 3000, 4160]
    else:
        widths = [9360 // col_count] * col_count
    set_table_width(table, widths)
    set_table_borders(table)

    header_cells = table.rows[0].cells
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, text in enumerate(headers):
        cell = header_cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, TABLE_HEADER)
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(text)
        run.bold = True
        run.font.color.rgb = DARK_BLUE

    for row_values in body:
        row_cells = table.add_row().cells
        for idx, text in enumerate(row_values):
            cell = row_cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.15
            run = para.add_run(text)
            run.font.size = Pt(10)

    doc.add_paragraph()


def split_table_row(line):
    """Split one pipe-delimited Markdown table row into trimmed cell text."""
    return [part.strip() for part in line.strip().strip("|").split("|")]


def build_doc():
    """Parse the proposal Markdown and write the fully formatted DOCX."""
    doc = Document()
    configure_page(doc)
    configure_styles(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].lstrip("#").strip()
    add_title_block(doc, title)

    table_buffer = []

    def flush_table():
        """Render and clear a contiguous block of buffered Markdown rows."""
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    # Parse only the Markdown constructs used by this proposal. Contiguous
    # table rows are buffered because the Word table must be created as one
    # object; all other blocks can be emitted immediately.
    for raw in lines[1:]:
        line = raw.rstrip()
        if not line:
            flush_table()
            continue

        if line.startswith("|"):
            parts = split_table_row(line)
            if all(re.fullmatch(r":?-{3,}:?", p) for p in parts):
                table_buffer.append(["---"] * len(parts))
            else:
                table_buffer.append(parts)
            continue

        flush_table()

        # Map the source's lightweight structure to native Word styles so the
        # resulting file remains editable and navigation-friendly.
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        elif line.startswith("- "):
            item = line[2:].strip()
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(item)
        elif line.endswith(":") and len(line) < 60:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(line)
            run.bold = True
            run.font.color.rgb = DARK_BLUE
        else:
            doc.add_paragraph(line)

    flush_table()

    # Populate document metadata separately from visible page content.
    doc.core_properties.title = "Research Proposal: Evaluating Explainability in VLMs for Construction Safety"
    doc.core_properties.subject = "Research proposal"
    doc.core_properties.keywords = "construction safety, VLM, XAI, explainability, Florence-2"
    doc.core_properties.author = ""
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
